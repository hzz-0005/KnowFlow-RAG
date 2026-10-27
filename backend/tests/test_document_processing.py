from io import BytesIO

import pytest
from docx import Document as WordDocument
from knowflow.app.document_processing import (
    ParsedSection,
    UnsupportedDocumentError,
    parse_document,
    split_sections,
)
from pypdf import PdfWriter


def test_markdown_parser_preserves_headings() -> None:
    parsed = parse_document("guide.md", "# 登录\n\n请先申请 API Key。\n".encode())

    assert [(part.section, part.text) for part in parsed] == [
        ("登录", "# 登录\n\n请先申请 API Key。"),
    ]


def test_markdown_parser_groups_content_by_heading_and_preserves_fenced_code() -> None:
    source = """# 数据库

第一段说明。
第二行仍属于同一段落。

```md
# 这不是标题
```

## 向量索引

使用余弦距离查询。
"""

    parsed = parse_document("guide.md", source.encode())

    assert [part.section for part in parsed] == ["数据库", "数据库 / 向量索引"]
    assert "第二行仍属于同一段落" in parsed[0].text
    assert "# 这不是标题" in parsed[0].text
    assert "使用余弦距离查询" in parsed[1].text


def test_markdown_chunk_count_follows_sections_and_content_size() -> None:
    source = "# A\n\n" + "A 段落说明。\n\n" * 8 + "# B\n\nB 段落说明。"

    chunks = split_sections(parse_document("guide.md", source.encode()))

    assert len(chunks) == 2
    assert [chunk.section for chunk in chunks] == ["A", "B"]
    assert all("段落说明" in chunk.content for chunk in chunks)


def test_splitter_keeps_provenance_and_bounds_long_text() -> None:
    chunks = split_sections([ParsedSection("第一段。" + "接口配置说明，" * 220, page_number=3)])

    assert len(chunks) > 1
    assert all(chunk.page_number == 3 for chunk in chunks)
    assert all(len(chunk.content) <= 1000 for chunk in chunks)
    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
    assert any("第一段" in chunk.content for chunk in chunks)


def test_docx_parser_reads_headings_paragraphs_and_tables() -> None:
    document = WordDocument()
    document.add_heading("数据库连接", level=1)
    document.add_paragraph("服务通过环境变量读取连接信息。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "参数"
    table.rows[0].cells[1].text = "示例值"
    document.add_heading("检索", level=1)
    document.add_paragraph("查询使用向量索引。")
    output = BytesIO()
    document.save(output)

    parsed = parse_document("runbook.docx", output.getvalue())

    assert any(part.section == "数据库连接" for part in parsed)
    assert any("环境变量" in part.text for part in parsed)
    assert "参数 | 示例值" in parsed[0].text
    assert "参数 | 示例值" not in parsed[1].text
    assert parsed[1].section == "检索"


def test_scanned_pdf_is_reported_instead_of_silently_indexed() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)

    with pytest.raises(ValueError, match="OCR"):
        parse_document("scan.pdf", output.getvalue())


def test_unsupported_file_type_has_actionable_error() -> None:
    with pytest.raises(UnsupportedDocumentError, match="TXT"):
        parse_document("archive.zip", b"some bytes")
