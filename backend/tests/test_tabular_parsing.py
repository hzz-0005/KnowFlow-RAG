from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from defusedxml.ElementTree import fromstring
from knowflow.app import tabular_parsing as tables
from knowflow.app.config import Settings
from knowflow.app.document_processing import parse_document
from knowflow.app.index_text import chunk_index_text
from knowflow.app.parse_artifacts import fingerprint, load_artifact, save_artifact
from knowflow.app.structured_parsing import chunk_parsed_document, native_parse
from knowflow.app.token_budget import get_token_counter
from openpyxl import Workbook


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "库存"
    sheet.append(["商品", "数量", "价格", "合计"])
    sheet.append(["001", 0, None, "=B2*C2"])
    sheet.append(["002", 3, 0.125, "=B3*C3"])
    sheet["B3"].number_format = "000000"
    sheet["C3"].number_format = "0.0%"
    sheet["D3"] = "=0"
    sheet.append([date(2026, 10, 23), "空字符串缓存", "错误缓存", "缺失缓存"])
    sheet["B4"] = '=IF(1,"","")'
    sheet["C4"] = "=1/0"
    sheet.row_dimensions[3].hidden = True
    hidden = book.create_sheet("归档")
    hidden.sheet_state = "hidden"
    hidden.append(["指标", "本期", "下期"])
    hidden.merge_cells("B1:C1")
    hidden.append(["名称", "实际", "预计"])
    hidden.append(["备用", 8, 9])
    raw = BytesIO()
    book.save(raw)
    # Test-only XML fixtures model a spreadsheet application's saved formula caches.
    from xml.etree.ElementTree import tostring

    out = BytesIO()
    with ZipFile(raw) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = fromstring(data)
                for cell in xml.iter(tables.NS + "c"):
                    address = cell.get("r")
                    if address in {"D3", "B4", "C4"}:
                        cell.set("t", {"D3": "n", "B4": "str", "C4": "e"}[address])
                        cell.find(tables.NS + "v").text = {"D3": "0", "B4": "", "C4": "#DIV/0!"}[
                            address
                        ]
                data = tostring(xml)
            target.writestr(info, data)
    return out.getvalue()


def cells(parsed, block=0):
    return {c["coordinate"]: c for c in parsed.blocks[block].structure["table_cells"]}


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "gb18030"])
def test_csv_encodings_literal_ids_quoted_newlines_and_blank_not_zero(encoding):
    parsed = tables.parse_tabular(
        "test.csv", '编号,数量,备注\r\n001,0,"跨行\r\n备注"\r\n002,,=1+1\r\n'.encode(encoding)
    )
    grid = cells(parsed)
    assert grid["A2"]["raw_value"] == "001"
    assert grid["B2"]["raw_value"] == "0"
    assert grid["B3"]["raw_value"] == ""
    assert grid["C3"]["formula"] is None
    assert grid["C3"]["raw_value"] == "=1+1"
    assert grid["C2"]["source_record"] == 2
    assert grid["C2"]["source_line_end"] == 3
    assert parsed.sections[1].page_number is None
    assert parsed.sections[1].source_spans[0]["cell_range"] == "A2:C2"


def test_csv_separator_directive_preserves_blank_boundaries_and_ragged_rows():
    parsed = tables.parse_tabular("test.csv", b"sep=;\r\n\r\nname;value\r\na;1\r\nb\r\n\r\nc;2\r\n")
    assert len(parsed.blocks) == 2
    assert cells(parsed)["A2"]["raw_value"] == "name"
    assert cells(parsed)["A2"]["source_line_end"] == 3
    assert parsed.report["tabular"]["ragged_records"] == [4]
    assert "B4" not in cells(parsed)
    assert cells(parsed, 1)["A6"]["raw_value"] == "c"


def test_xlsx_preserves_multisheet_merge_format_dates_and_formula_states():
    parsed = tables.parse_tabular("test.xlsx", workbook_bytes())
    grid = cells(parsed)
    assert grid["A2"]["raw_value"] == "001"
    assert grid["B2"]["raw_value"] == 0
    assert grid["C2"]["raw_value"] is None
    assert grid["D2"]["formula_cache_status"] == "missing"
    assert "未知" in grid["D2"]["display_value"]
    assert grid["D3"]["cached_value"] == 0
    assert grid["D3"]["formula_cache_status"] == "available_unverified"
    assert grid["B4"]["cached_value"] == ""
    assert grid["B4"]["formula_cache_status"] == "available_unverified"
    assert grid["C4"]["formula_cache_status"] == "cached_error"
    assert "错误" in grid["C4"]["display_value"]
    assert grid["B3"]["display_value"] == "000003"
    assert grid["B3"]["raw_value"] == 3
    assert grid["B3"]["row_hidden"] is True
    assert grid["C3"]["display_value"] == "12.5%"
    assert grid["C3"]["raw_value"] == 0.125
    assert grid["A4"]["raw_value"].startswith("2026-10-23")
    archived = cells(parsed, 1)
    assert archived["C1"]["raw_value"] is None
    assert archived["C1"]["merged_from"] == "B1"
    assert archived["C1"]["display_value"] == "本期"
    assert parsed.report["sheets"][1]["sheet_state"] == "hidden"
    assert parsed.report["tabular"]["formula_cells_missing_cache"] == 1
    assert parsed.report["tabular"]["formula_cells_cached_error"] == 1
    assert parsed.report["tabular"]["formula_cells_cached"] == 2
    assert parsed.report["page_count"] is None
    assert len(parsed.sections) == 7


@pytest.mark.parametrize("payload", [b"", b"\x00", b'a,b\n"unclosed', b"a,b\n\xff\xff"])
def test_invalid_csv_fails_without_replacing_or_silently_dropping_data(payload):
    with pytest.raises(ValueError):
        tables.parse_tabular("bad.csv", payload)


def test_limits_fail_explicitly_before_merge_expansion(monkeypatch):
    monkeypatch.setattr(tables, "MAX_ROWS", 2)
    with pytest.raises(ValueError, match="上限"):
        tables.parse_tabular("bad.csv", b"a,b\nx,y\nz,w")
    book = Workbook()
    book.active["A1"] = "title"
    book.active.merge_cells("A1:A3")
    raw = BytesIO()
    book.save(raw)
    # The preflight must reject before the library is allowed to load/expand merges.
    monkeypatch.setattr(tables, "load_workbook", lambda *a, **kw: pytest.fail("late limit check"))
    with pytest.raises(ValueError, match="上限"):
        tables.parse_tabular("bad.xlsx", raw.getvalue())


def test_xlsx_invalid_and_zip_expansion_limit(monkeypatch):
    with pytest.raises(ValueError, match="损坏或加密"):
        tables.parse_tabular("bad.xlsx", b"not a zip")
    data = workbook_bytes()
    monkeypatch.setattr(tables, "MAX_EXPANDED_BYTES", 10)
    with pytest.raises(ValueError, match="解压后"):
        tables.parse_tabular("large.xlsx", data)


def test_structure_chunks_cover_all_rows_keep_sheets_separate_and_obey_input_budget():
    settings = Settings(chunk_max_tokens=256, chunk_target_tokens=160, chunk_overlap_tokens=12)
    parsed = native_parse("test.xlsx", workbook_bytes(), settings)
    chunks = chunk_parsed_document(parsed, settings)
    counter = get_token_counter(settings.chunk_tokenizer_path)
    rows = set()
    for chunk in chunks:
        assert (
            counter.count(chunk_index_text(chunk.section, chunk.content))
            + settings.chunk_token_reserve
            <= 256
        )
        assert len({s["sheet_name"] for s in chunk.source_spans}) == 1
        assert chunk.page_number is None
        rows.update((s["sheet_name"], s["row_start"]) for s in chunk.source_spans)
    assert rows == {("库存", i) for i in range(1, 5)} | {("归档", i) for i in range(1, 4)}
    assert len(parse_document("t.csv", b"id,value\n001,0")) == 2


def test_tabular_cache_roundtrip_and_pdf_fingerprint_backward_compatibility(tmp_path):
    import hashlib
    import json

    settings = Settings(upload_dir=str(tmp_path))
    parsed = tables.parse_tabular("t.xlsx", workbook_bytes())
    save_artifact("file.xlsx", "t.xlsx", "sha", parsed, settings)
    loaded = load_artifact("file.xlsx", "t.xlsx", "sha", settings)
    assert loaded.blocks == parsed.blocks
    assert loaded.sections == parsed.sections
    expected = {
        "schema": 1,
        "source_sha256": "sha",
        "extension": ".pdf",
        "parser": settings.pdf_parser,
        "parser_revision": settings.parse_artifact_revision,
        "endpoint": settings.docling_base_url if settings.pdf_parser == "docling" else None,
        "max_pages": settings.pdf_max_pages,
    }
    assert (
        fingerprint("old.pdf", "sha", settings)
        == hashlib.sha256(
            json.dumps(expected, ensure_ascii=False, sort_keys=True).encode()
        ).hexdigest()
    )
