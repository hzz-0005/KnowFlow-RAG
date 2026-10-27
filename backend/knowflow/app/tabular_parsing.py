"""Bounded CSV/XLSX ingestion with literal cells, formula state and stable coordinates."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

VERSION = "tabular-cells-v1"
MAX_CELLS = 200_000
MAX_COLUMNS = 256
MAX_ROWS = 50_000
MAX_SHEETS = 64
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_TEXT = 5_000_000
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and not Decimal(str(value)).is_finite():
        raise ValueError("表格包含无法表示的非有限数值。")
    return value


def _display(value, number_format):
    value = _json_value(value)
    if value is None:
        return "[空白，非零]"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if re.fullmatch(r"0{2,20}", number_format) and value == int(value):
            return f"{int(value):0{len(number_format)}d}"
        match = re.fullmatch(r"0(?:\.(0{1,10}))?%", number_format)
        if match:
            decimals = len(match[1] or "")
            return f"{Decimal(str(value)) * 100:.{decimals}f}%"
    return str(value) if not isinstance(value, bool) else ("true" if value else "false")


def _csv_grid(data):
    encoding = "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError:
        encoding = "gb18030"
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as exc:
            raise ValueError(
                "CSV 编码无法读取；请使用 UTF-8、带 BOM 的 UTF-16 或 GB18030。"
            ) from exc
    if "\x00" in text:
        raise ValueError("CSV 含 NUL 字符，请检查文件类型和编码。")
    if len(text) > MAX_TEXT:
        raise ValueError("CSV 文字量超过处理上限，请按数据范围拆分。")
    delimiter, directive = None, None
    first = text.splitlines()[0] if text else ""
    if re.fullmatch(r"sep=[,;\t|]", first, re.I):
        delimiter, directive = first[-1], first
        text = text[len(first) :]
        text = text[2:] if text.startswith("\r\n") else text[1:]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows, consumed = [], 0
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        for row_number, values in enumerate(reader, 1):
            consumed += len(values)
            if row_number > MAX_ROWS or len(values) > MAX_COLUMNS or consumed > MAX_CELLS:
                raise ValueError("CSV 行列或单元格数量超过处理上限，请拆分文件。")
            rows.append(
                [
                    {
                        "coordinate": f"{get_column_letter(col)}{row_number}",
                        "row": row_number,
                        "column": col,
                        "raw_value": value,
                        "display_value": value,
                        "data_type": "string",
                        "number_format": None,
                        "formula": None,
                        "cached_value": None,
                        "formula_cache_status": "not_formula",
                        "source_record": row_number,
                        "source_line_end": reader.line_num + bool(directive),
                    }
                    for col, value in enumerate(values, 1)
                ]
            )
    except csv.Error as exc:
        raise ValueError(f"CSV 记录格式错误，位于物理行 {reader.line_num} 附近。") from exc
    width = max(map(len, rows), default=0)
    return (
        [
            {
                "sheet_name": "CSV",
                "sheet_state": "visible",
                "rows": rows,
                "width": width,
                "merged_ranges": [],
            }
        ],
        {
            "encoding": encoding,
            "delimiter": delimiter,
            "separator_directive": directive,
            "ragged_records": [i + 1 for i, row in enumerate(rows) if row and len(row) != width],
        },
        [],
    )


def _xlsx_grids(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            info = archive.infolist()
            if len(info) > 10_000 or sum(i.file_size for i in info) > MAX_EXPANDED_BYTES:
                raise ValueError("XLSX 解压后超过处理上限，请拆分工作簿。")
            names = archive.namelist()
            # The XML store distinguishes a missing cache, an empty string and numeric zero.
            raw_sheets = {}
            scanned_area = 0
            for name in names:
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    xml = fromstring(archive.read(name))
                    # Check before openpyxl expands merged areas into Python objects.
                    max_row, max_col = 0, 0
                    areas = [c.get("r") for c in xml.iter(NS + "c")]
                    areas += [c.get("ref") for c in xml.iter(NS + "mergeCell")]
                    for area in filter(None, areas):
                        _, _, right, bottom = range_boundaries(area)
                        max_row, max_col = max(max_row, bottom), max(max_col, right)
                    scanned_area += max_row * max_col
                    if max_row > MAX_ROWS or max_col > MAX_COLUMNS or scanned_area > MAX_CELLS:
                        raise ValueError(
                            "XLSX 已用区域超过行列/单元格处理上限，请清理远端格式或拆分。"
                        )
                    raw_sheets[name] = {
                        c.attrib["r"]: {
                            "stored_value": c.findtext(NS + "v"),
                            "stored_type": c.get("t", "n"),
                            "formula_text": c.findtext(NS + "f"),
                            "formula_kind": (
                                c.find(NS + "f").get("t", "normal")
                                if c.find(NS + "f") is not None
                                else None
                            ),
                        }
                        for c in xml.iter(NS + "c")
                        if "r" in c.attrib
                    }
                    if sum(map(len, raw_sheets.values())) > MAX_CELLS:
                        raise ValueError("XLSX 单元格数量超过处理上限，请拆分工作簿。")
            omitted = {
                "charts": sum(
                    n.startswith("xl/charts/chart") and n.endswith(".xml") for n in names
                ),
                "media": sum(n.startswith("xl/media/") for n in names),
            }
        book = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 XLSX 工作簿，请检查文件是否损坏或加密。") from exc
    grids, scanned, text_size, unknown, available, errors = [], 0, 0, [], [], []
    try:
        if len(book.worksheets) > MAX_SHEETS:
            raise ValueError("XLSX 工作表数量超过处理上限，请拆分工作簿。")
        for sheet in book.worksheets:
            scanned += sheet.max_row * sheet.max_column
            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS or scanned > MAX_CELLS:
                raise ValueError("XLSX 已用区域超过行列/单元格处理上限，请清理远端格式或拆分。")
            # Worksheet paths are obtained from workbook relationships, not ordinal filenames.
            xml_path = cached[sheet.title]._worksheet_path
            stored = raw_sheets.get(xml_path, {})
            merges, inherited = [], {}
            for area in sheet.merged_cells.ranges:
                merges.append(str(area))
                for row in range(area.min_row, area.max_row + 1):
                    for col in range(area.min_col, area.max_col + 1):
                        inherited[row, col] = sheet.cell(area.min_row, area.min_col).coordinate
            cache_values = {
                c.coordinate: c.value
                for row in cached[sheet.title].iter_rows()
                for c in row
                if c.value is not None
            }
            rows = []
            for cells in sheet.iter_rows():
                row_values = []
                for cell in cells:
                    anchor = inherited.get((cell.row, cell.column))
                    value_cell = sheet[anchor] if anchor else cell
                    value = value_cell.value
                    raw_cell = stored.get(cell.coordinate, {})
                    raw = (
                        _json_value(cell.value)
                        if cell.data_type != "f" or isinstance(cell.value, str)
                        else "=" + (raw_cell.get("formula_text") or "[复杂公式，未计算]")
                    )
                    formula = (
                        value if value_cell.data_type == "f" and isinstance(value, str) else None
                    )
                    source = stored.get(value_cell.coordinate, {})
                    if value_cell.data_type == "f" and not formula:
                        formula = "=" + (source.get("formula_text") or "[复杂公式，未计算]")
                    cached_value = None
                    cache_status = "not_formula"
                    if formula:
                        cache = cache_values.get(value_cell.coordinate)
                        has_cache = (
                            source.get("stored_value") not in (None, "")
                            or source.get("stored_value") == ""
                            and source.get("stored_type") == "str"
                        )
                        if has_cache and source.get("stored_type") == "str" and cache is None:
                            cache = ""
                        cached_value = _json_value(cache)
                        cache_status = "available_unverified" if has_cache else "missing"
                        location = f"{sheet.title}!{value_cell.coordinate}"
                        if has_cache and source.get("stored_type") == "e":
                            cache_status = "cached_error"
                            errors.append(location)
                        else:
                            (available if has_cache else unknown).append(location)
                        shown = (
                            f"[公式缓存错误：{cache}，未计算]"
                            if cache_status == "cached_error"
                            else _display(cache, value_cell.number_format) + "（公式缓存，未核算）"
                            if has_cache
                            else "[公式缓存未知，未计算]"
                        )
                    elif value_cell.data_type == "e":
                        shown = f"[单元格错误：{value}]"
                    else:
                        shown = _display(value, value_cell.number_format)
                    text_size += len(shown) + len(formula or "")
                    if text_size > MAX_TEXT:
                        raise ValueError("XLSX 文字量超过处理上限，请按数据范围拆分。")
                    row_values.append(
                        {
                            "coordinate": cell.coordinate,
                            "row": cell.row,
                            "column": cell.column,
                            "raw_value": raw,
                            "display_value": shown,
                            "raw_data_type": cell.data_type,
                            "data_type": value_cell.data_type,
                            "number_format": value_cell.number_format,
                            "merged_from": anchor if anchor != cell.coordinate else None,
                            "formula": formula,
                            "cached_value": cached_value,
                            "formula_cache_status": cache_status,
                            "row_hidden": bool(sheet.row_dimensions[cell.row].hidden),
                            **source,
                        }
                    )
                rows.append(row_values)
            grids.append(
                {
                    "sheet_name": sheet.title,
                    "sheet_state": sheet.sheet_state,
                    "rows": rows,
                    "width": sheet.max_column,
                    "merged_ranges": merges,
                }
            )
    finally:
        book.close()
        cached.close()
    warnings = []
    if unknown:
        warnings.append(f"{len(set(unknown))} 个公式单元格没有保存的计算结果；已标未知，没有求值。")
    if available:
        warnings.append("公式缓存可能过期；仅保留文件中的缓存，没有重新计算或验证。")
    if errors:
        warnings.append(f"{len(set(errors))} 个公式单元格缓存为错误值，未当作有效数值。")
    if omitted["charts"] or omitted["media"]:
        warnings.append("工作簿含图表/图片：本次仅提取单元格，尚未解释这些视觉内容。")
    return (
        grids,
        {
            "formula_cells_missing_cache": len(set(unknown)),
            "formula_cells_cached": len(set(available)),
            "formula_cells_cached_error": len(set(errors)),
            "missing_cache_examples": list(dict.fromkeys(unknown))[:50],
            "omitted_visuals": omitted,
        },
        warnings,
    )


def _occupied(row):
    return any(
        c.get("raw_value") not in (None, "") or c.get("merged_from") or c.get("formula")
        for c in row
    )


def parse_tabular(filename, data):
    # Lazy imports keep the shared parser-neutral dataclasses free of a module cycle.
    from knowflow.app.document_processing import ParsedSection
    from knowflow.app.structured_parsing import Block, ParsedDocument

    suffix = Path(filename).suffix.lower()
    grids, metadata, warnings = _csv_grid(data) if suffix == ".csv" else _xlsx_grids(data)
    blocks, sections, sheets = [], [], []
    for sheet_index, grid in enumerate(grids):
        nonempty = sum(_occupied(r) for r in grid["rows"])
        sheets.append(
            {k: grid[k] for k in ("sheet_name", "sheet_state", "merged_ranges")}
            | {
                "row_count": len(grid["rows"]),
                "nonempty_rows": nonempty,
                "column_count": grid["width"],
            }
        )
        # Blank record boundaries partition separate tables; no rows are capped or sampled.
        regions, active = [], []
        for row in grid["rows"]:
            if _occupied(row):
                active.append(row)
            elif active:
                regions.append(active)
                active = []
        if active:
            regions.append(active)
        for region_index, rows in enumerate(regions):
            key = f"sheet-{sheet_index}-table-{region_index}"
            first, last = rows[0][0]["row"], rows[-1][0]["row"]
            width = max(len(r) for r in rows)
            caption = (
                f"工作表：{grid['sheet_name']} · 范围 A{first}:{get_column_letter(width)}{last}"
            )
            # Candidate labels supply context, never remove the candidate header record itself.
            header = next(
                (r for r in rows[:8] if sum(c.get("raw_value") not in (None, "") for c in r) >= 2),
                rows[0],
            )
            header_rows = [r for r in rows if r[0]["row"] == header[0]["row"]]
            if any(c.get("merged_from") for c in header) and len(rows) > 1:
                child = next((r for r in rows if r[0]["row"] == header[0]["row"] + 1), None)
                if child:
                    header_rows.append(child)
            labels = {}
            for col in range(1, width + 1):
                values = [
                    str(r[col - 1]["display_value"])
                    for r in header_rows
                    if col <= len(r) and r[col - 1]["display_value"] not in ("", "[空白，非零]")
                ]
                labels[col] = " / ".join(dict.fromkeys(values))
            table_text, row_sections, table_cells = [], [], []
            for row in rows:
                rnum = row[0]["row"]
                parts = []
                for cell in row:
                    label = labels.get(cell["column"], "")
                    if len(label) > 160:
                        label = label[:160] + "…（列名节选）"
                    shown = cell["display_value"]
                    if cell.get("formula"):
                        shown += f"；公式={cell['formula']}"
                    if cell.get("merged_from"):
                        shown += f"（合并自 {cell['merged_from']}，本格原值为空）"
                    parts.append(
                        f"[{cell['coordinate']}] {label} = {json.dumps(shown, ensure_ascii=False)}"
                    )
                    table_cells.append(
                        {
                            **cell,
                            "text": str(cell["display_value"]),
                            "column_header": rnum in {r[0]["row"] for r in header_rows},
                            "start_row_offset_idx": rnum - first,
                            "end_row_offset_idx": rnum - first + 1,
                            "start_col_offset_idx": cell["column"] - 1,
                            "end_col_offset_idx": cell["column"],
                        }
                    )
                text = caption + "\n" + " | ".join(parts)
                span = {
                    "block_key": key,
                    "source_type": "spreadsheet",
                    "sheet_name": grid["sheet_name"],
                    "sheet_state": grid["sheet_state"],
                    "row_start": rnum,
                    "row_end": rnum,
                    "cell_range": f"A{rnum}:{get_column_letter(len(row))}{rnum}",
                    "row_key": str(
                        next(
                            (
                                c["display_value"]
                                for c in row
                                if c.get("raw_value") not in (None, "")
                            ),
                            "",
                        )
                    ),
                    "cell_keys": [v for v in labels.values() if v],
                    "header_inference": "candidate_not_verified",
                }
                span["header_cell_ranges"] = [
                    f"A{r[0]['row']}:{get_column_letter(len(r))}{r[0]['row']}" for r in header_rows
                ]
                table_text.append(text)
                row_sections.append(ParsedSection(text, section=caption, source_spans=[span]))
            spans = [s.source_spans[0] for s in row_sections]
            blocks.append(
                Block(
                    key,
                    "table",
                    "\n\n".join(table_text),
                    caption,
                    source_spans=spans,
                    structure={
                        "num_rows": len(rows),
                        "num_cols": width,
                        "table_cells": table_cells,
                        "captions": [caption],
                        "sheet_name": grid["sheet_name"],
                        "merged_ranges": grid["merged_ranges"],
                        "header_inference": "candidate_not_verified",
                        "source_format": suffix[1:],
                    },
                )
            )
            sections.extend(row_sections)
    if not sections:
        raise ValueError("表格中没有可索引的单元格内容。")
    warnings.append("列名为自动推断的候选上下文，所有原行均保留；请以工作表坐标和原值为准。")
    if metadata.get("ragged_records"):
        warnings.append("CSV 存在列数不一致的记录；已按实际列保留，没有补零或丢弃。")
    return ParsedDocument(
        sections,
        blocks,
        {
            "schema_version": 1,
            "parser": "tabular",
            "parser_version": VERSION,
            "status": "needs_review" if warnings else "parsed",
            "page_count": None,
            "pages": [],
            "block_count": len(blocks),
            "table_count": len(blocks),
            "picture_count": 0,
            "source_format": suffix[1:],
            "sheets": sheets,
            "sheet_count": len(sheets),
            "ocr_requested": False,
            "warnings": warnings,
            "tabular": metadata,
            "note": (
                "按工作表/行列定位，无固定页码；公式不执行，空白不作零。"
                "解析完成不代表答案正确。"
            ),
        },
    )
