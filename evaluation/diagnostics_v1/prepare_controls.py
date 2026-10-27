"""Create synthetic controls once; refuse to overwrite a frozen corpus."""

import hashlib
import json
import shutil
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]


def main():
    corpus = HERE / "corpus"
    corpus.mkdir(exist_ok=False)
    (corpus / "flow.md").write_text(
        "# 合成控制：资料流程\n\n"
        + "\n\n".join(
            f"## 第 {i} 节\n"
            + (f"步骤 {i} 的输入应保留来源编号与说明，后续处理不得改变原始值。" * 20)
            for i in range(1, 4)
        ),
        encoding="utf-8",
    )
    doc = Document()
    doc.add_heading("Synthetic operational controls", 0)
    doc.add_paragraph("Record QX-17 links to the receipt below. These are synthetic test facts.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Record"
    table.rows[0].cells[1].text = "Receipt"
    table.add_row().cells[0].text = "QX-17"
    table.rows[1].cells[1].text = "R-082"
    doc.add_paragraph("The table is followed by a note: preserve the original identifier.")
    doc.save(corpus / "flow.docx")
    files = [
        {"filename": name, "source": "synthetic_authored_control"}
        for name in ("flow.md", "flow.docx")
    ]
    for name in ("service-control.csv", "inventory-control.xlsx"):
        source = ROOT / "evaluation/tabular_v1" / name
        shutil.copyfile(source, corpus / name)
        files.append({"filename": name, "source": source.relative_to(ROOT).as_posix()})
    for row in files:
        row["sha256"] = hashlib.sha256((corpus / row["filename"]).read_bytes()).hexdigest()
    (HERE / "corpus-manifest.json").write_text(
        json.dumps(
            {
                "role": "synthetic_software_controls_not_semantic_gold",
                "files": files,
            },
            ensure_ascii=True,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps({"synthetic_inputs": len(files), "overwritten": False}))


if __name__ == "__main__":
    main()
